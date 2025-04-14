# type: ignore
import pathlib
from openpyxl import Workbook

ROOT_FOLDER = pathlib.Path(__file__).parent
WORKBOOK_FOLDER = ROOT_FOLDER / "workbook.xlsx"

sheet_name = "Minha planilha"
workbook = Workbook()
workbook.create_sheet(sheet_name, 0)  # type: ignore

worksheet = workbook[sheet_name]


# criando os cabeçalhos
worksheet.cell(row=1, column=1, value="Nome")
worksheet.cell(row=1, column=2, value="idade")
worksheet.cell(row=1, column=3, value="nota")

students = [
    # nome      idade nota
    ["João", 14, 5.5],
    ["Maria", 13, 9.7],
    ["Luiz", 15, 8.8],
    ["Alberto", 16, 10],
]

# def insert_data(list_students: list, worksheet_: Workbook):
#     for i, student in enumerate(list_students, start=2):
#         for j, student_data in enumerate(student, start=1):
#             worksheet_.cell(row=i, column=j, value=student_data)

# insere os dados no arquivo excel workbook
# insert_data(students, worksheet)  # type: ignore


for i in students:
    worksheet.append(i)

# cria o arquivo excel workbook
workbook.save(WORKBOOK_FOLDER)
